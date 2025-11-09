"""
Resume Parser & Scoring Agent — Single-file implementation (step-by-step)

What this file contains (section-by-section):
1) Setup & requirements (how to install deps and set env vars)
2) Config & helper utilities
3) Google Sheets helpers (read/write master/detail sheets using gspread)
4) Resume extraction utilities (PDF/DOCX/text) — best-effort local extraction
5) Candidate data enrichment helpers (basic heuristics for education, experience, location)
6) Rule-based scorer (configurable scoring rules loaded from master sheet or a local dict)
7) Notifier (SMTP example for Gmail; placeholders for Gmail API/Twilio)
8) Scheduler (APScheduler example) and follow-up workflow
9) CLI entrypoint + examples showing how to run

Notes & disclaimers:
- Replace placeholders with your Google service account credentials and sheet IDs.
- For production use: switch to the Gmail API (or Twilio) with OAuth, secure secrets, add retries and full logging.
- This single-file layout is intended for clarity; you may split into modules for maintainability.

Dependencies (pip):
- gspread
- google-auth
- pandas
- python-docx
- pdfminer.six
- apscheduler
- python-dotenv
- fuzzywuzzy (optional, for fuzzy skill matching) or rapidfuzz
- requests (if you fetch resume URLs)

Install example:
    python -m pip install gspread google-auth pandas python-docx pdfminer.six apscheduler python-dotenv rapidfuzz requests

Environment variables (use .env file or export):
- GOOGLE_SERVICE_ACCOUNT_JSON: Path to service account JSON file (or set up gspread as you prefer)
- MASTER_SHEET_ID: Google Sheet ID for master_sheet
- DETAIL_SHEET_ID: Google Sheet ID for detail_sheet
- SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASSWORD (or use OAuth/credentials)
- FROM_EMAIL

Run example:
    python resume_parser_agent.py --run-once
    python resume_parser_agent.py --schedule

"""

# ---------------------------
# SECTION 0: Imports
# ---------------------------
import os
import re
import io
import json
import argparse
import logging
from datetime import datetime, timedelta
from typing import Dict, List, Optional, Tuple

import pandas as pd

# Google Sheets
import gspread
from google.oauth2.service_account import Credentials

# Resume parsing
from pdfminer.high_level import extract_text as extract_pdf_text
from docx import Document

# Notifications & scheduling
import smtplib
from email.message import EmailMessage
from apscheduler.schedulers.background import BackgroundScheduler

# Optional fuzzy matching
try:
    from rapidfuzz import fuzz
except Exception:
    fuzz = None

# ---------------------------
# SECTION 1: Logging & config
# ---------------------------
logging.basicConfig(level=logging.INFO, format='%(asctime)s %(levelname)s: %(message)s')
logger = logging.getLogger('resume-parser-agent')

# Load environment variables
from dotenv import load_dotenv
load_dotenv()

GOOGLE_SERVICE_ACCOUNT_JSON = os.getenv('GOOGLE_SERVICE_ACCOUNT_JSON')
MASTER_SHEET_ID = os.getenv('MASTER_SHEET_ID')
DETAIL_SHEET_ID = os.getenv('DETAIL_SHEET_ID')
SMTP_HOST = os.getenv('SMTP_HOST', 'smtp.gmail.com')
SMTP_PORT = int(os.getenv('SMTP_PORT', 587))
SMTP_USER = os.getenv('SMTP_USER')
SMTP_PASSWORD = os.getenv('SMTP_PASSWORD')
FROM_EMAIL = os.getenv('FROM_EMAIL', SMTP_USER)

# Default follow-ups (8 messages) and default schedule offsets (days after initial)
DEFAULT_FOLLOWUPS = [
    (1, "Thanks for applying — we received your resume."),
    (2, "Reminder: We are reviewing your application."),
    (4, "Update: Your application is under consideration."),
    (7, "Interview scheduling — next steps."),
    (10, "Final reminder — keep an eye on your inbox."),
    (14, "Status update from recruitment team."),
    (18, "Last call for confirmation."),
    (24, "Closing the application process. Thank you.")
]

# ---------------------------
# SECTION 2: Google Sheets Helpers
# ---------------------------
SCOPES = [
    'https://www.googleapis.com/auth/spreadsheets',
    'https://www.googleapis.com/auth/drive'
]


def get_gspread_client(json_path: Optional[str] = None) -> gspread.Client:
    """Create a gspread client using service account JSON."""
    if json_path is None:
        json_path = GOOGLE_SERVICE_ACCOUNT_JSON
    if json_path is None:
        raise ValueError("GOOGLE_SERVICE_ACCOUNT_JSON must be set (path to service account json)")
    creds = Credentials.from_service_account_file(json_path, scopes=SCOPES)
    client = gspread.authorize(creds)
    return client


class SheetsWrapper:
    def __init__(self, client: gspread.Client, master_sheet_id: str, detail_sheet_id: str):
        self.client = client
        self.master_ss = client.open_by_key(master_sheet_id)
        self.detail_ss = client.open_by_key(detail_sheet_id)

    def read_master(self, sheet_name: str = 'master') -> pd.DataFrame:
        sh = self.master_ss.worksheet(sheet_name)
        data = sh.get_all_records()
        return pd.DataFrame(data)

    def read_detail(self, sheet_name: str = 'detail') -> pd.DataFrame:
        sh = self.detail_ss.worksheet(sheet_name)
        data = sh.get_all_records()
        return pd.DataFrame(data)

    def update_detail_from_df(self, df: pd.DataFrame, sheet_name: str = 'detail'):
        sh = self.detail_ss.worksheet(sheet_name)
        # Replace entire sheet values
        sh.clear()
        sh.update([df.columns.values.tolist()] + df.values.tolist())

    def append_detail_row(self, row: Dict, sheet_name: str = 'detail'):
        sh = self.detail_ss.worksheet(sheet_name)
        values = [row.get(c, '') for c in sh.row_values(1)]
        sh.append_row(values)

# ---------------------------
# SECTION 3: Resume extraction utilities
# ---------------------------

def extract_text_from_pdf(path_or_bytes: bytes | str) -> str:
    """Extract text from PDF file path or bytes. If bytes provided, use io.BytesIO."""
    try:
        if isinstance(path_or_bytes, (bytes, bytearray)):
            text = extract_pdf_text(io.BytesIO(path_or_bytes))
        else:
            text = extract_pdf_text(path_or_bytes)
        return text
    except Exception as e:
        logger.exception('PDF parse failed: %s', e)
        return ''


def extract_text_from_docx(path: str) -> str:
    try:
        doc = Document(path)
        return '\n'.join(p.text for p in doc.paragraphs)
    except Exception as e:
        logger.exception('DOCX parse failed: %s', e)
        return ''


def fetch_resume_text_from_url(url: str) -> str:
    """If detail_sheet stores a resume URL, fetch it and extract text. Uses requests."""
    import requests
    try:
        r = requests.get(url, timeout=20)
        r.raise_for_status()
        ctype = r.headers.get('Content-Type','')
        if 'pdf' in ctype or url.lower().endswith('.pdf'):
            return extract_text_from_pdf(r.content)
        elif 'word' in ctype or url.lower().endswith('.docx'):
            # Save to temp and parse
            with open('/tmp/temp_resume.docx','wb') as f:
                f.write(r.content)
            return extract_text_from_docx('/tmp/temp_resume.docx')
        else:
            return r.text
    except Exception as e:
        logger.exception('Failed fetching resume url: %s', e)
        return ''

# ---------------------------
# SECTION 4: Candidate enrichment heuristics
# ---------------------------

EDUCATION_PATTERNS = {
    'btech': r"\bB\.?\s?Tech\b|Bachelor\s+of\s+Technology|BTech\b",
    'bsc': r"\bB\.?\s?Sc\b|Bachelor\s+of\s+Science",
    'mtech': r"\bM\.?\s?Tech\b|Master\s+of\s+Technology",
    'msc': r"\bM\.?\s?Sc\b|Master\s+of\s+Science",
    'mba': r"\bMBA\b|Master\s+of\s+Business\s+Administration",
}

LOCATION_KEYWORDS = [
    'Hyderabad', 'Bengaluru', 'Bangalore', 'Pune', 'Chennai', 'Mumbai', 'Delhi'
]


def detect_education(text: str) -> List[str]:
    found = []
    for key, pat in EDUCATION_PATTERNS.items():
        if re.search(pat, text, flags=re.I):
            found.append(key)
    return found


def detect_location(text: str) -> Optional[str]:
    for loc in LOCATION_KEYWORDS:
        if re.search(r'\b' + re.escape(loc) + r'\b', text, flags=re.I):
            return loc
    return None


def estimate_experience_months(text: str) -> int:
    """Very rough heuristic: looks for phrases like 'X years', 'Y months'."""
    years = re.findall(r"(\d+)\s+years?", text, flags=re.I)
    months = re.findall(r"(\d+)\s+months?", text, flags=re.I)
    total = 0
    if years:
        total += sum(int(y) * 12 for y in years)
    if months:
        total += sum(int(m) for m in months)
    # fallback: look for internship durations 'intern for 6 months'
    return total

# ---------------------------
# SECTION 5: Scoring engine
# ---------------------------

DEFAULT_SCORING = {
    'education': {
        'btech': 10,
        'mtech': 12,
        'mba': 8
    },
    'top_tier_college': 15,
    'experience_in_months_threshold': { 'months': 5, 'points': 15 },
    'location': { 'Hyderabad': 15 },
    'profile_tagline': 10,
    'resume_quality': 15,
}


def score_candidate(candidate: Dict, scoring_config: Dict = None) -> Dict:
    """Compute rule-based score. candidate is a dict with keys: education_text, location, experience_months, tagline, resume_text"""
    if scoring_config is None:
        scoring_config = DEFAULT_SCORING
    score_breakdown = {}
    total = 0

    # Education
    educs = candidate.get('education_text','')
    educ_found = detect_education(educs)
    educ_points = 0
    for e in educ_found:
        educ_points = max(educ_points, scoring_config.get('education', {}).get(e, 0))
    score_breakdown['education'] = educ_points
    total += educ_points

    # Top tier school: placeholder: check if school in candidate['college'] matches a list
    top_tier = 0
    college_name = candidate.get('college','') or ''
    top_tier_list = scoring_config.get('top_tier_list', ['IIT', 'NIT', 'BITS'])
    for t in top_tier_list:
        if t.lower() in college_name.lower():
            top_tier = scoring_config.get('top_tier_college', 0)
            break
    score_breakdown['top_tier_college'] = top_tier
    total += top_tier

    # Experience
    exp_months = int(candidate.get('experience_months') or 0)
    exp_cfg = scoring_config.get('experience_in_months_threshold', {})
    exp_points = exp_cfg.get('points', 0) if exp_months >= exp_cfg.get('months', 9999) else 0
    score_breakdown['experience'] = exp_points
    total += exp_points

    # Location
    location_points = scoring_config.get('location', {}).get(candidate.get('location'), 0)
    score_breakdown['location'] = location_points
    total += location_points

    # Profile tagline: simple heuristic (presence)
    tagline_pts = scoring_config.get('profile_tagline', 0) if candidate.get('tagline') else 0
    score_breakdown['tagline'] = tagline_pts
    total += tagline_pts

    # Resume quality: heuristic based on length
    resume_text = candidate.get('resume_text','') or ''
    word_count = len(re.findall(r"\w+", resume_text))
    resume_pts = min(scoring_config.get('resume_quality', 0), int(word_count / 50))  # example: 50 words = 1 point
    score_breakdown['resume_quality'] = resume_pts
    total += resume_pts

    # Cap total at 100
    total = min(total, 100)
    score_breakdown['total'] = total
    return score_breakdown

# ---------------------------
# SECTION 6: Notifier (SMTP simple)
# ---------------------------

def send_email_smtp(to_email: str, subject: str, body: str, from_email: str = FROM_EMAIL) -> bool:
    if not SMTP_USER or not SMTP_PASSWORD:
        logger.warning('SMTP credentials not set; skipping send_email')
        return False
    try:
        msg = EmailMessage()
        msg['Subject'] = subject
        msg['From'] = from_email
        msg['To'] = to_email
        msg.set_content(body)
        with smtplib.SMTP(SMTP_HOST, SMTP_PORT) as s:
            s.starttls()
            s.login(SMTP_USER, SMTP_PASSWORD)
            s.send_message(msg)
        logger.info('Sent mail to %s', to_email)
        return True
    except Exception as e:
        logger.exception('Failed to send email: %s', e)
        return False

# ---------------------------
# SECTION 7: Follow-up scheduler and workflow
# ---------------------------

scheduler = BackgroundScheduler()


def schedule_followups_for_candidate(candidate_row: Dict, start_date: Optional[datetime] = None):
    """Schedule follow-ups via APScheduler. For demonstration, we schedule in the background.
    In real deployment, use persistent job store (e.g., Redis, DB) or Cron.
    """
    if start_date is None:
        start_date = datetime.now()
    candidate_email = candidate_row.get('email')
    candidate_id = candidate_row.get('id') or candidate_row.get('email')

    for idx, (offset_days, message) in enumerate(DEFAULT_FOLLOWUPS, start=1):
        run_date = start_date + timedelta(days=offset_days-1)
        job_id = f"followup-{candidate_id}-{idx}"

        def make_job(to_email, subject, body):
            def job():
                send_email_smtp(to_email, subject, body)
            return job

        scheduler.add_job(make_job(candidate_email, f"Follow-up #{idx}", message), 'date', run_date=run_date, id=job_id)
        logger.info('Scheduled followup %s for %s at %s', idx, candidate_email, run_date)

# ---------------------------
# SECTION 8: End-to-end processing
# ---------------------------

def process_master_and_details(client_wrapper: SheetsWrapper, scoring_config: Dict = None, run_followups: bool = False):
    master_df = client_wrapper.read_master()
    detail_df = client_wrapper.read_detail()

    # Example: iterate master rows (job postings), for each posting, load candidate urls from the detail sheet link column
    # For simplicity, we'll assume detail_df already has candidate rows with at least: email, resume_url, resume_text

    updated_rows = []
    for idx, row in detail_df.iterrows():
        candidate = dict(row)
        # If resume_text missing, try to fetch from URL or local path
        if not candidate.get('resume_text') and candidate.get('resume_url'):
            candidate['resume_text'] = fetch_resume_text_from_url(candidate['resume_url'])
        elif not candidate.get('resume_text') and candidate.get('resume_path'):
            rp = candidate['resume_path']
            if rp.lower().endswith('.pdf'):
                candidate['resume_text'] = extract_text_from_pdf(rp)
            elif rp.lower().endswith('.docx'):
                candidate['resume_text'] = extract_text_from_docx(rp)

        # Enrich
        candidate['education_text'] = '\n'.join(detect_education(candidate.get('resume_text','')))
        candidate['location'] = candidate.get('location') or detect_location(candidate.get('resume_text',''))
        candidate['experience_months'] = candidate.get('experience_months') or estimate_experience_months(candidate.get('resume_text',''))

        # Score
        score_br = score_candidate(candidate, scoring_config=scoring_config)
        candidate.update({ 'score_total': score_br['total'], 'score_breakdown': json.dumps(score_br) })

        # Add scheduling (if needed)
        if run_followups and candidate.get('email'):
            schedule_followups_for_candidate(candidate)

        updated_rows.append(candidate)

    # Convert back to DataFrame and write
    updated_df = pd.DataFrame(updated_rows)

    # ensure columns are strings/lists flattened
    client_wrapper.update_detail_from_df(updated_df)
    logger.info('Updated detail sheet with %d candidates', len(updated_df))

# ---------------------------
# SECTION 9: CLI and entrypoint
# ---------------------------

def main():
    parser = argparse.ArgumentParser(description='Resume Parser & Scoring Agent')
    parser.add_argument('--run-once', action='store_true', help='Run the ETL once and exit')
    parser.add_argument('--schedule', action='store_true', help='Start scheduler (keeps running)')
    parser.add_argument('--master-id', type=str, default=MASTER_SHEET_ID)
    parser.add_argument('--detail-id', type=str, default=DETAIL_SHEET_ID)
    args = parser.parse_args()

    if not args.master_id or not args.detail_id:
        logger.error('Provide MASTER_SHEET_ID and DETAIL_SHEET_ID either as env vars or CLI args')
        return

    client = get_gspread_client()
    wrapper = SheetsWrapper(client, args.master_id, args.detail_id)

    if args.run_once:
        process_master_and_details(wrapper, scoring_config=DEFAULT_SCORING, run_followups=True)
        logger.info('Run-once completed')
        return

    if args.schedule:
        # Start scheduler and run job periodically (e.g., every day at 02:00) — here: every 12 hours for demo
        scheduler.add_job(lambda: process_master_and_details(wrapper, scoring_config=DEFAULT_SCORING, run_followups=True), 'interval', hours=12, id='daily-run')
        scheduler.start()
        logger.info('Scheduler started — running in background. Press Ctrl+C to exit.')
        try:
            # Keep main thread alive
            import time
            while True:
                time.sleep(60)
        except (KeyboardInterrupt, SystemExit):
            scheduler.shutdown()
            logger.info('Scheduler stopped')

if __name__ == '__main__':
    main()
